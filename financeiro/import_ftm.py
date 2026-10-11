"""Reads the original workbook; produces a private, traceable import, never committed."""
import argparse
import json
import time
import unicodedata
from decimal import Decimal
from pathlib import Path
import openpyxl

def text(v):return unicodedata.normalize('NFKD',str(v or '')).encode('ascii','ignore').decode().upper().strip()
def value(v):
    if isinstance(v,(int,float,Decimal)) and not isinstance(v,bool) and v>0:return int(Decimal(str(v)).quantize(Decimal('.01'))*100)
    return None
def category(description):
    s=text(description)
    for cat,terms in {'Moradia':['ALUGUEL','CASA','LUZ','NET'],'Saúde':['SAUDE','FARMACIA'],'Educação':['ESCOLA','POS','INGLES','FORMAT'],'Transporte':['CARRO','GASOLINA'],'Impostos':['SEFAZ'],'Alimentação':['FEIRA','COMIDA']}.items():
        if any(t in s for t in terms):return cat
    return 'Outros'

def extract(path):
    w=openpyxl.load_workbook(path,read_only=True,data_only=False)
    months=['JANEIRO','FEVEREIRO','MARCO','ABRIL','MAIO','JUNHO','JULHO','AGOSTO','SETEMBRO','OUTUBRO','NOVEMBRO','DEZEMBRO'];records=[]
    for sheet in w:
        normalized=text(sheet.title)
        if normalized not in months:continue
        period=f'2026-{months.index(normalized)+1:02}'
        offset=1 if text(sheet.cell(1,1).value)=='JHC' else 0
        desc_col=2+offset; amount_col=3+offset; status_col=4+offset
        for row in range(1,sheet.max_row+1):
            income=value(sheet.cell(row,1+offset).value)
            source_name=sheet.cell(row,1).value if offset else None
            if income and text(source_name)!='TOTAL':
                name=str(source_name).strip() if offset else f'Receita {row} — origem a confirmar'
                records.append(dict(description=name,kind='income',category='Receitas',month=period,amount=income,paid=0,review=1,source=f'FTM2026:{normalized}:{sheet.cell(row,1+offset).coordinate}',notes='Confirmar origem e se a receita foi recebida. A planilha não informa data de recebimento.',created=int(time.time())))
            description=sheet.cell(row,desc_col).value; total=value(sheet.cell(row,amount_col).value)
            if description and total and text(description) not in ('TOTAL','DESCRICAO','DESPESAS'):
                card='CART' in text(description)
                state=sheet.cell(row,status_col).value
                note=f'Importado de {sheet.title}, linha {row}. Situação original: {state if state is not None else "em branco"}. '
                note+='Não há vencimento informado. '
                if isinstance(state,(int,float)):note+='O número na situação pode indicar pagamento parcial; confirmar antes de usar.'
                if card:note+='Total de fatura histórica; não adicionar as mesmas compras novamente. Conferir os dias de fechamento e vencimento do cartão.'
                records.append(dict(description=str(description).strip(),kind='invoice' if card else 'expense',card_name=str(description).strip() if card else None,category='Cartões — histórico' if card else category(description),month=period,amount=total,paid=total if text(state)=='OK' else 0,review=1,source=f'FTM2026:{normalized}:{sheet.cell(row,amount_col).coordinate}',notes=note,created=int(time.time())))
        # Independent blocks on the right change headings/positions by month.
        headers={};description_columns={}
        for row in range(1,sheet.max_row+1):
            for col in range(7+offset,sheet.max_column+1):
                cell=sheet.cell(row,col);v=cell.value;s=text(v)
                if isinstance(v,str) and not v.startswith('=') and s in ('COMIDA FORA','FARMACIA','VARIAVEIS','GASOLINA','COMIDA','TURISMO'):
                    headers[col]=str(v).strip();continue
                total=value(v)
                if total and col in headers:
                    adjacent=sheet.cell(row,col+1).value if col<sheet.max_column else None
                    description=str(adjacent).strip() if isinstance(adjacent,str) and not adjacent.startswith('=') else headers[col]
                    records.append(dict(description=description,kind='expense',category='Outros' if text(headers[col])=='VARIAVEIS' else category(headers[col]),month=period,amount=total,paid=0,review=1,source=f'FTM2026:{normalized}:{cell.coordinate}',notes='Bloco adicional da planilha. Confirmar se já faz parte da fatura de cartão antes de incluir no total.',created=int(time.time())))
    return records

if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('source');p.add_argument('output');a=p.parse_args();records=extract(a.source)
    target=Path(a.output);target.parent.mkdir(parents=True,exist_ok=True);target.write_text(json.dumps(records,ensure_ascii=False),encoding='utf-8')
    print(json.dumps({'records':len(records),'months':len({r['month'] for r in records}),'review':sum(r['review'] for r in records),'invoices':sum(r['kind']=='invoice' for r in records)}))
